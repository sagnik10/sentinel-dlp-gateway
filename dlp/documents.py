"""Bounded local extraction and privacy-preserving document reports.

Document content is data: macros, formulas, links and embedded instructions are
never executed. Unsupported or incomplete extraction blocks conservatively.
"""
import hashlib
import html
import io
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from .utils import analyze_prompt

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED = 80 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_UNITS = 10000


def extract(data):
    """Return detected type, located text units, and extraction limitations."""
    units, issues = [], []
    total = 0

    def add(location, text):
        nonlocal total
        text = str(text or '').strip()
        if not text:
            return
        total += len(text)
        if total > MAX_TEXT or len(units) >= MAX_UNITS:
            raise ValueError('Extraction limit exceeded')
        units.append((location, text))

    if not data or len(data) > MAX_BYTES:
        raise ValueError('Empty file or file size limit exceeded')
    if data.startswith(b'%PDF-'):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ValueError('Encrypted PDF requires a decrypted copy')
        if len(reader.pages) > 500:
            raise ValueError('Page limit exceeded')
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            add(f'page {number}', text)
            if not text.strip() or len(page.images):
                issues.append(f'page {number}: image or missing text requires OCR review')
            for ai, annotation in enumerate(page.get('/Annots', []), 1):
                item = annotation.get_object()
                add(f'page {number}, annotation {ai}', item.get('/Contents', ''))
                if item.get('/A') or item.get('/AA'):
                    issues.append(f'page {number}: active annotation requires separate review')
        for fi, field in enumerate((reader.get_fields() or {}).values(), 1):
            add(f'form field {fi}', field.get('/V', ''))
        for mi, value in enumerate((reader.metadata or {}).values(), 1):
            add(f'metadata field {mi}', value)
        root = reader.trailer['/Root']
        if root.get('/OpenAction') or root.get('/AA') or '/JavaScript' in root.get('/Names', {}):
            issues.append('Active PDF content requires separate review')
        if reader.attachments:
            issues.append('Embedded PDF attachments require separate inspection')
        return 'pdf', units, issues
    if not data.startswith(b'PK'):
        raise ValueError('Unsupported format; use XLSX, DOCX or PDF (convert legacy XLS/DOC first)')
    with ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_EXPANDED:
            raise ValueError('Archive expansion limit exceeded')
        names = set(archive.namelist())
        if any('embeddings/' in n or 'vbaProject' in n or '/media/' in n for n in names):
            issues.append('Images, macros or embedded objects require separate inspection')
        if 'xl/workbook.xml' in names:
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for si, sheet in enumerate(book, 1):
                    sheet.reset_dimensions()
                    for ri, row in enumerate(sheet.iter_rows(), 1):
                        if ri > MAX_UNITS:
                            raise ValueError('Row limit exceeded')
                        values = [str(c.value) for c in row if c.value is not None]
                        if any(c.data_type == 'f' for c in row):
                            issues.append(f'sheet {si}, row {ri}: formula results are not evaluated')
                        add(f'sheet {si}, row {ri}', ' | '.join(values))
            finally:
                book.close()
            # Comments and ancillary XML text can also carry sensitive values.
            from defusedxml.ElementTree import fromstring
            for ni, name in enumerate(sorted(names), 1):
                if name.endswith('.xml') and ('comments' in name.lower() or name.startswith('docProps/')):
                    root = fromstring(archive.read(name))
                    add(f'package part {ni}', ' '.join(root.itertext()))
            return 'xlsx', units, issues
        if 'word/document.xml' in names:
            from defusedxml.ElementTree import fromstring
            ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            for ni, name in enumerate(sorted(names), 1):
                if name.startswith('word/') and name.endswith('.xml'):
                    root = fromstring(archive.read(name))
                    for index, para in enumerate(root.iter(ns + 'p'), 1):
                        add(f'word part {ni}, paragraph {index}', ''.join(para.itertext()))
                elif name.startswith('docProps/') and name.endswith('.xml'):
                    add(f'metadata part {ni}', ' '.join(fromstring(archive.read(name)).itertext()))
            return 'docx', units, issues
    raise ValueError('Unrecognized Office package')


def inspect_document(data):
    report = {'schema_version': 1, 'generated_at': datetime.now(timezone.utc).isoformat(),
              'sha256': hashlib.sha256(data).hexdigest() if len(data) <= MAX_BYTES else None,
              'bytes_read': len(data), 'input_over_limit': len(data) > MAX_BYTES,
              'decision': 'BLOCKED', 'format': 'unknown', 'complete': False,
              'scope': 'Extractable text only; no OCR or embedded-object execution',
              'issues': [], 'findings': [], 'forwarded': False}
    try:
        kind, units, issues = extract(data)
        report.update(format=kind, issues=sorted(set(issues)))
        if not units:
            report['issues'].append('No extractable text')
        # Overlap preserves matches across text chunk boundaries. A separate
        # stream pass catches secrets split across adjacent rows or paragraphs.
        for location, text in units:
            for start in range(0, len(text), 7500):
                result = analyze_prompt(text[start:start + 8000])
                report['findings'].append({'location': location, 'offset': start,
                                           'characters': len(text[start:start + 8000]), **result})
        combined = '\n'.join(t for _, t in units)
        for start in range(0, len(combined), 7500):
            result = analyze_prompt(combined[start:start + 8000])
            if result['decision'] == 'BLOCKED':
                report['findings'].append({'location': 'cross-unit text', 'offset': start, **result})
        if any('Security validation unavailable' in f['safe_detection_summary'] for f in report['findings']):
            report['issues'].append('One or more detectors were unavailable')
        report['complete'] = bool(units) and not report['issues']
        if report['complete'] and all(f['decision'] == 'ALLOWED' for f in report['findings']):
            report['decision'] = 'ALLOWED'
    except Exception:
        report['issues'].append('Extraction failed: unsupported, encrypted, malformed, over-limit file or missing parser dependency')
    report['counts'] = dict(Counter(f['decision'] for f in report['findings']))
    return report


def write_reports(report, output):
    """Write JSON and escaped standalone HTML without raw document excerpts."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.with_suffix('.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    escaped = html.escape(json.dumps(report, indent=2))
    findings = report.get('findings', report.get('results', []))
    rows = ''.join('<tr><td>' + '</td><td>'.join(html.escape(str(v)) for v in (
        f.get('location', ''), f.get('decision', ''), f.get('expected', ''),
        ', '.join(f.get('matched_policy_codes', [])),
        '; '.join(f.get('safe_detection_summary', [])))) + '</td></tr>' for f in findings)
    summary = (f"Decision: {report['decision']} | Complete: {report.get('complete', False)}"
               if 'decision' in report else f"Matched: {report['matched']} / {report['labeled']} labeled prompts")
    page = ('<!doctype html><html lang="en"><meta charset="utf-8"><title>DLP inspection report</title>'
            '<style>body{font:16px system-ui;max-width:1100px;margin:40px auto;padding:24px;'
            'background:#f4f7fb;color:#17243b}pre{white-space:pre-wrap;overflow-wrap:anywhere;'
            'background:white;padding:24px;border-radius:12px}h1{color:#164e63}'
            'table{width:100%;border-collapse:collapse;background:white}td,th{padding:10px;text-align:left;'
            'border-bottom:1px solid #ddd;overflow-wrap:anywhere}th{background:#164e63;color:white}</style>'
            '<h1>DLP inspection report</h1><p>Detailed local decisions. Document text is omitted.</p>'
            f'<h2>{html.escape(summary)}</h2><table><tr><th>Location</th><th>Decision</th>'
            f'<th>Expected</th><th>Policies</th><th>Reasons</th></tr>{rows}</table>'
            f'<details><summary>Full report and extraction limitations</summary><pre>{escaped}</pre></details></html>')
    output.with_suffix('.html').write_text(page, encoding='utf-8')
