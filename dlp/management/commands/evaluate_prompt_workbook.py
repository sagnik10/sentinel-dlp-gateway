import hashlib
from collections import Counter
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.test.utils import override_settings
from dlp.documents import MAX_BYTES, extract, write_reports
from dlp.utils import analyze_prompt


def expected_decision(value):
    value = str(value or '').strip().casefold()
    if value in {'allow', 'allowed', 'safe'}:
        return 'ALLOWED'
    if value in {'blocked', 'block', 'block / anonymize', 'block/anonymize', 'block/redact', 'block/flag'}:
        return 'BLOCKED'
    return None


class Command(BaseCommand):
    help = 'Evaluate labeled XLSX prompts without training on them or echoing their content.'

    def add_arguments(self, parser):
        parser.add_argument('workbook', type=Path)
        parser.add_argument('--output', type=Path, default=Path('reports/prompt_evaluation'))
        parser.add_argument('--rules-only', action='store_true', help='Explicitly evaluate deterministic rules without optional models')

    def handle(self, *args, **options):
        if options['rules_only']:
            with override_settings(DLP_ENABLE_LOCAL_CLASSIFIER=False, DLP_ENABLE_PRESIDIO=False, DLP_GUARD_MODELS=[]):
                return self.evaluate(options)
        return self.evaluate(options)

    def evaluate(self, options):
        from openpyxl import load_workbook
        import io
        try:
            with options['workbook'].open('rb') as stream:
                data = stream.read(MAX_BYTES + 1)
            if extract(data)[0] != 'xlsx':
                raise ValueError()
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        except Exception:
            raise CommandError('Workbook cannot be safely parsed') from None
        rows, issues, cached_results = [], [], {}
        try:
            for si, sheet in enumerate(book, 1):
                header = None
                for ri, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                    keys = [str(c or '').strip().casefold().replace(' ', '_') for c in cells]
                    prompt_key = next((k for k in ('prompt_text', 'prompt') if k in keys), None)
                    if prompt_key and 'expected_action' in keys:
                        header = (keys.index(prompt_key), keys.index('expected_action'))
                        continue
                    if header is None and ri == 1 and len(cells) == 2 and expected_decision(cells[1]):
                        header = (0, 1)
                    if not any(c is not None for c in cells):
                        continue
                    location = f'sheet {si}, row {ri}'
                    if header is None:
                        issues.append({'location': location, 'issue': 'No recognized prompt/expected header'})
                        continue
                    pi, ei = header
                    prompt = cells[pi] if pi < len(cells) else None
                    expected = expected_decision(cells[ei] if ei < len(cells) else None)
                    if not isinstance(prompt, str) or not prompt.strip():
                        issues.append({'location': location, 'issue': 'Missing prompt'})
                        continue
                    digest = hashlib.sha256(prompt.encode()).hexdigest()
                    if digest not in cached_results:
                        cached_results[digest] = analyze_prompt(prompt)
                    result = cached_results[digest]
                    rows.append({'location': location, 'sha256': digest,
                                 'expected': expected, 'matches_expected': result['decision'] == expected if expected else None,
                                 **result})
                    if expected is None:
                        issues.append({'location': location, 'issue': 'Missing or ambiguous expected action; excluded from accuracy'})
        finally:
            book.close()
        labeled = [r for r in rows if r['expected']]
        matched = sum(r['matches_expected'] for r in labeled)
        labels_by_hash = {}
        for row in labeled:
            labels_by_hash.setdefault(row['sha256'], set()).add(row['expected'])
        conflicts = {key for key, labels in labels_by_hash.items() if len(labels) > 1}
        report = {'mode': 'prompt_benchmark', 'sha256': hashlib.sha256(data).hexdigest(),
                  'rules_only': options['rules_only'],
                  'mapping': 'Block / Anonymize maps to BLOCKED; ambiguous labels are excluded.',
                  'rows': len(rows), 'labeled': len(labeled), 'matched': matched,
                  'accuracy': matched / len(labeled) if labeled else None,
                  'decisions': dict(Counter(r['decision'] for r in rows)),
                  'false_allowed': sum(r['expected'] == 'BLOCKED' and r['decision'] == 'ALLOWED' for r in rows),
                  'false_blocked': sum(r['expected'] == 'ALLOWED' and r['decision'] == 'BLOCKED' for r in rows),
                  'detector_errors': sum('Security validation unavailable' in r['safe_detection_summary'] for r in rows),
                  'unique_prompts': len({r['sha256'] for r in rows}),
                  'conflicting_label_groups': len(conflicts),
                  'conflicting_label_locations': [r['location'] for r in rows if r['sha256'] in conflicts],
                  'issues': issues, 'results': rows}
        write_reports(report, options['output'])
        self.stdout.write(f'Evaluated {len(rows)} prompts; matched {matched}/{len(labeled)} labeled rows; {len(issues)} data issues.')
